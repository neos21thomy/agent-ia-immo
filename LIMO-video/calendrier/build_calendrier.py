"""Calendrier de publication 4 semaines (Instagram + Facebook) : légendes prêtes à copier-coller."""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF", size=11)
HF = PatternFill("solid", fgColor="2A1F6B")
B = Font(name=F, size=10)
BB = Font(name=F, size=10, bold=True)
T = Font(name=F, bold=True, size=16, color="2A1F6B")
thin = Side(style="thin", color="D5D0EE")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical="top")
WEEK = [PatternFill("solid", fgColor=c) for c in ("FFFFFF", "F4F1FD", "FFFFFF", "F4F1FD")]

H1 = "#conseillerimmobilier #agentimmobilier #immobilier #mandataireimmobilier #limo"
H2 = "#immobilierfrance #negociateurimmobilier #prospectionimmobilier #productivite #limo"
H3 = "#agentimmo #immo #estimationimmobiliere #mandatexclusif #outilsimmo #limo"
CTA = "👉 Essai gratuit 14 jours, sans carte bancaire : lien en bio."
N = "\n"

# (fichier, thème, légende, commentaire épinglé, hashtags)
POSTS = [
    ("premium-pack/9x16/prem-02-trois-erreurs.mp4", "3 erreurs qui font perdre des mandats",
     f"3 erreurs qui te font perdre des mandats… sans que tu t'en rendes compte 👇{N}{N}1️⃣ Tu rappelles trop tard après l'estimation{N}2️⃣ Tes relances sont dans ta tête{N}3️⃣ Tes soirées partent dans l'administratif{N}{N}Laquelle te parle le plus ? Dis-le en commentaire 👇{N}{N}{CTA}",
     "Tu veux la démo ? Commente « LIMO » et je t'envoie le lien en MP 📩", H1),
    ("premium-pack/9x16/prem-03-dimanche-soir.mp4", "POV dimanche soir",
     f"Dimanche, 22 h. Toi, tu ressors tes notes pour savoir qui rappeler… 😩{N}Ton collègue, lui, regarde un film 🍿{N}{N}Son secret : Limo prépare sa semaine (relances, rendez-vous, dossiers).{N}{N}Identifie le collègue qui bosse encore le dimanche 😅{N}{N}{CTA}",
     "Et toi, tu fais quoi le dimanche soir ? 😅", H2),
    ("hooks/LIMO-hook-05-ton-concurrent-musique.mp4", "Ton concurrent",
     f"Pendant que tu hésites, ton concurrent rappelle tes vendeurs. 📞{N}{N}Le mandat ne se gagne pas pendant l'estimation : il se gagne au suivi.{N}{N}{CTA}",
     "Tu relances tes vendeurs combien de fois après une estimation ?", H3),
    ("premium-pack/9x16/prem-04-le-calcul.mp4", "Le calcul (temps d'admin)",
     f"Fais le calcul ⏱️{N}1 h d'administratif par jour = plus de 200 h par an.{N}{N}200 h, c'est des dizaines de rendez-vous vendeurs que tu ne fais pas.{N}Limo s'occupe des relances, comptes rendus, annonces et suivi des dossiers.{N}{N}Dès 49 €/mois. {CTA}",
     "Combien de temps tu passes sur l'admin par jour ? Sois honnête 😄", H1),
    ("mascotte/LIMO-mascotte-01-salut-agent-immo-musique.mp4", "Présentation de la mascotte",
     f"Salut conseiller immo 👋 C'est moi, le robot de Limo.{N}{N}Mon job : te soulager de tout ce qui n'est pas de la vente. Relances, fiches, annonces, dossiers… je m'en occupe.{N}{N}{CTA}",
     "Donne-moi un prénom en commentaire 🤖", H2),
    ("cine/9x16/LIMO-cine-04-prestige-9x16.mp4", "Biens de prestige",
     f"Tu vends des biens d'exception. Ton suivi doit l'être aussi ✨{N}{N}Estimation, annonce soignée, relances au bon moment : Limo, le bras droit du conseiller immo.{N}{N}{CTA}",
     "C'est quoi le plus beau bien que tu as vendu ? 🏡", H3),
    ("premium-pack/9x16/prem-01-regarde-ca.mp4", "Démo : regarde ça",
     f"Conseiller immo, regarde ça 👀{N}{N}Tu demandes « rappelle-moi le bien de M. et Mme T. »… et tout ressort : documents, propriétaires, estimation, diagnostics.{N}Pendant ce temps, Limo relance tes acquéreurs et ton diagnostiqueur.{N}{N}{CTA}",
     "Tu veux voir l'app en vrai ? Commente « DÉMO » 📩", H1),
    ("voix/9x16/voix-01-julien.mp4", "L'histoire de Julien",
     f"Julien a fait une super estimation. Puis il a oublié de rappeler.{N}Quelques jours plus tard, un autre panneau était sur la maison. 💔{N}{N}Le mandat ne se perd pas pendant l'estimation. Il se perd après.{N}{N}Même les meilleurs conseillers oublient parfois. Limo, jamais.{N}{N}{CTA}",
     "Ça t'est déjà arrivé ? Raconte 👇", H2),
    ("pubs-30s/9x16/LIMO-pub30-01-mandats-perdus-9x16.mp4", "Mandats perdus",
     f"Combien de mandats as-tu perdus cette année… juste par manque de suivi ? 🤔{N}{N}Limo relance tes vendeurs au bon moment et garde tous tes dossiers à jour.{N}{N}{CTA}",
     "Un chiffre en commentaire, sans jugement 😉", H3),
    ("lifestyle/9x16/LIMO-life-02-pov-collegue-9x16.mp4", "POV : le collègue",
     f"POV : ton collègue a l'air moins débordé que toi… et rentre plus tôt 😅{N}{N}Son secret ? Limo, le bras droit du conseiller immo.{N}{N}Tag le collègue qui a besoin de voir ça 👇{N}{N}{CTA}",
     "Tag un collègue débordé 👇", H1),
    ("mascotte/LIMO-mascotte-03-3h-du-matin-musique.mp4", "3 h du matin",
     f"3 h du matin. Tu te réveilles en sursaut : « J'ai rappelé Mme Martin ?! » 😱{N}{N}Avec Limo, tu dors tranquille : les relances sont suivies pour toi.{N}{N}{CTA}",
     "Ça t'est déjà arrivé ? 🙋", H2),
    ("cine/9x16/LIMO-cine-05-chalet-9x16.mp4", "Chalet / montagne",
     f"Un chalet, une vue, un acquéreur qui attend la bonne annonce 🏔️{N}{N}Avec Limo, l'annonce est rédigée en 2 minutes et tes acquéreurs compatibles sont relancés.{N}{N}{CTA}",
     "Montagne ou bord de mer ? 🏔️🌊", H3),
    ("premium-pack/9x16/prem-05-une-journee.mp4", "Une journée avec Limo",
     f"8 h : priorités prêtes. 10 h : compte rendu dicté, fiche créée. 14 h : annonce rédigée. 18 h : diagnostiqueur relancé, notaire servi.{N}{N}Et toi ? Tu as passé ta journée avec tes clients 🏡{N}{N}{CTA}",
     "C'est quoi la tâche que tu détestes le plus dans ta journée ?", H1),
    ("agents/9x16/LIMO-agents-les-12-agents-9x16.mp4", "Les outils de Limo",
     f"Tout ce que Limo fait pour toi, en une vidéo 🧰{N}{N}Relances, estimations, annonces, dossiers, avis, suivi acquéreurs… le bras droit du conseiller immo.{N}{N}Lequel tu utiliserais en premier ?{N}{N}{CTA}",
     "Réponds par un chiffre : quel outil en premier ?", H2),
    ("hooks/LIMO-hook-04-red-flags-musique.mp4", "Red flags",
     f"Les red flags d'une journée de conseiller immo 🚩{N}{N}Tu en coches combien ? Dis-le en commentaire 👇{N}{N}{CTA}",
     "Ton pire red flag ? 🚩", H3),
    ("pubs-30s/9x16/LIMO-pub30-05-crm-dort-9x16.mp4", "Le CRM qui dort",
     f"Ton CRM est plein… mais il dort 😴{N}{N}Limo te dit qui relancer, quand, et prépare le message pour toi.{N}{N}{CTA}",
     "Combien de contacts dorment dans ton CRM ?", H1),
    ("mascotte/LIMO-mascotte-05-vrai-ou-faux-musique.mp4", "Vrai ou faux",
     f"Vrai ou faux ? Le robot de Limo répond à tes idées reçues 🤖{N}{N}Ta réponse en commentaire 👇{N}{N}{CTA}",
     "Vrai ou faux ? 👇", H2),
    ("signatures/9x16/sign-01-derriere-chaque-signature.mp4", "Derrière chaque signature (émotion)",
     f"Le jour de la signature, on oublie tout le reste 🤝{N}Les appels du soir, les relances, les compromis, les diagnostics…{N}{N}Ton métier, c'est ça. Limo s'occupe du reste, du mandat à la signature.{N}{N}{CTA}",
     "Ton plus beau souvenir de signature ? ❤️", H3),
    ("voix/9x16/voix-03-surcharge.mp4", "Surchargé ?",
     f"Conseiller immo, surchargé ? 😮‍💨{N}{N}Limo est fait pour ça.{N}{N}{CTA}",
     "Note ta charge de travail cette semaine sur 10 👇", H1),
    ("pubs-30s/9x16/LIMO-pub30-06-avis-google-9x16.mp4", "Avis Google",
     f"Tes avis Google, ta meilleure vitrine ⭐{N}{N}Limo t'aide à demander l'avis au bon moment et à y répondre.{N}{N}{CTA}",
     "Tu as combien d'avis Google ? ⭐", H2),
    ("hooks/LIMO-hook-03-dix-secondes-musique.mp4", "10 secondes",
     f"10 secondes pour comprendre Limo ⏱️{N}{N}{CTA}",
     "Tu as compris ? Dis-moi en 1 mot 👇", H3),
    ("lifestyle/9x16/LIMO-life-01-pendant-ce-temps-9x16.mp4", "Pendant ce temps",
     f"Toi, tu décroches un mandat depuis ta voiture.{N}Pendant ce temps, Limo envoie l'estimation à la vendeuse 📲{N}{N}{CTA}",
     "Tu travailles depuis ta voiture toi aussi ? 🚗", H1),
    ("mascotte/LIMO-mascotte-04-il-reagit-musique.mp4", "La mascotte réagit",
     f"Le robot de Limo réagit à ta journée de conseiller 😂{N}{N}{CTA}",
     "Envoie ça à un collègue 😂", H2),
    ("cine/9x16/LIMO-cine-01-manifeste-9x16.mp4", "Manifeste",
     f"Ton métier, c'est le terrain, l'humain, la signature.{N}Le reste ? Limo s'en occupe.{N}{N}Limo, le bras droit du conseiller immo.{N}{N}{CTA}",
     "Qu'est-ce que tu aimes le plus dans ton métier ? ❤️", H3),
]
STORIES = [
    "Repartage le Reel du jour + sondage « Tu relances tes vendeurs : de tête / dans un carnet / dans un logiciel ? »",
    "Repartage le Reel + question « Ta tâche la plus pénible de la semaine ? »",
    "Coulisses : capture de l'app Limo + « Tu veux tester ? Réponds OUI »",
    "Repartage le Reel + curseur emoji 😩→😎 « Ton niveau de charge cette semaine »",
    "Quiz « Combien de temps pour rédiger une annonce avec Limo ? 2 min / 20 min / 1 h »",
    "Repartage le Reel + lien vers l'essai gratuit (sticker lien)",
    "Repos ou story perso (terrain, visite, signature) : montre le côté humain",
]

wb = Workbook()
ws = wb.active
ws.title = "Calendrier"
ws["A1"] = "Calendrier de publication LIMO · 4 semaines · Instagram + Facebook"; ws["A1"].font = T
ws["A2"] = "Copie la légende (colonne G) et le commentaire épinglé (H). Vidéo : dossier LIMO-video/rendus/. Coche le statut une fois publié."
ws["A2"].font = Font(name=F, italic=True, size=10, color="6B4FE0")
cols = ["Sem.", "Jour", "Date", "Heure", "Format", "Vidéo à publier", "Légende (copier-coller)", "Commentaire à épingler", "Hashtags", "Story du jour", "Statut"]
for i, c in enumerate(cols, 1):
    cell = ws.cell(row=4, column=i, value=c); cell.font, cell.fill, cell.border = H, HF, BOX
    cell.alignment = Alignment(wrap_text=True, vertical="center")
for col, w in zip("ABCDEFGHIJK", (6, 10, 12, 8, 16, 44, 70, 34, 40, 40, 12)):
    ws.column_dimensions[col].width = w
JOURS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
start = datetime.date(2026, 10, 5)
dv = DataValidation(type="list", formula1='"À publier,Programmé,Publié"'); ws.add_data_validation(dv)
r, k = 5, 0
for d in range(28):
    day = start + datetime.timedelta(days=d)
    wd = day.weekday(); wk = d // 7
    if wd == 6:
        vals = [wk + 1, JOURS[wd], day, "", "Story seule", "—", "Pas de Reel le dimanche : récupère et réponds aux commentaires de la semaine.", "", "", STORIES[6], "À publier"]
    else:
        f, theme, cap, com, tags = POSTS[k]; k += 1
        hour = "10:00" if wd == 5 else ("12:15" if wd % 2 == 0 else "18:30")
        vals = [wk + 1, JOURS[wd], day, hour, "Reel (Insta + FB)", f"{f}\n→ {theme}", cap, com, tags, STORIES[wd], "À publier"]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v); c.font = B; c.alignment = WR; c.border = BOX; c.fill = WEEK[wk]
    ws.cell(row=r, column=3).number_format = "ddd dd/mm"
    ws.cell(row=r, column=11).font = Font(name=F, size=10, color="0000FF", bold=True)
    dv.add(f"K{r}")
    ws.row_dimensions[r].height = 170 if wd != 6 else 50
    r += 1
ws.conditional_formatting.add(f"K5:K{r-1}", CellIsRule(operator="equal", formula=['"Publié"'], fill=PatternFill("solid", fgColor="CDEFE9")))
ws.cell(row=r + 1, column=6, value="Publiés").font = BB
ws.cell(row=r + 1, column=7, value=f'=COUNTIF(K5:K{r-1},"Publié")&" / "&COUNTIF(E5:E{r-1},"Reel (Insta + FB)")&" Reels"').font = BB
ws.freeze_panes = "F5"

# onglet pubs
p = wb.create_sheet("Pubs Meta")
p["A1"] = "Calendrier des pubs (Meta Ads, 20 €/jour)"; p["A1"].font = T
hd = ["Période", "Action", "Vidéos", "Budget", "À surveiller"]
for i, c in enumerate(hd, 1):
    cell = p.cell(row=3, column=i, value=c); cell.font, cell.fill, cell.border = H, HF, BOX
rows = [
    ("Lun 5 → Dim 11 oct.", "Lancer la vague 1 (campagne Prospects, formulaire instantané)", "prem-02 Trois erreurs · prem-03 Dimanche soir · prem-04 Le calcul", "20 €/jour", "Ne rien toucher 5 jours. Rappeler chaque contact dans l'heure."),
    ("Lun 12 oct.", "Couper la pub la plus chère par contact", "—", "20 €/jour", "Coût par contact, taux de hook > 25 %, CTR > 1 %"),
    ("Lun 12 → Dim 18 oct.", "Vague 2 : ajouter 2 nouvelles vidéos", "prem-01 Regarde ça · prem-05 Une journée", "20 €/jour", "Comparer avec les 2 gagnantes de la vague 1"),
    ("Lun 19 oct.", "Relance : nouvelle campagne sur les personnes qui ont vu 50 % d'une vidéo", "sign-01 Signatures (si accord clients) · voix-01 Julien", "+5 €/jour", "Coût par démo réservée"),
    ("Lun 19 → Dim 1 nov.", "Augmenter de 20 % ce qui est rentable (voir onglet Pilotage du kit pub)", "Les 2 meilleures", "24-30 €/jour", "Coût par essai démarré < ~37 €"),
]
for i, row in enumerate(rows):
    for j, v in enumerate(row, 1):
        c = p.cell(row=4 + i, column=j, value=v); c.font = B; c.alignment = WR; c.border = BOX
for col, w in zip("ABCDE", (22, 50, 50, 14, 50)):
    p.column_dimensions[col].width = w

# onglet conseils
g = wb.create_sheet("Règles d'or")
g["A1"] = "Règles d'or pour publier"; g["A1"].font = T
tips = [
    "Publie la vidéo en Reel sur Instagram et coche « Partager aussi sur Facebook » : un seul geste pour les deux réseaux.",
    "Mets la miniature (dossier rendus/premium-pack/miniatures/) quand elle existe : c'est elle qu'on voit dans ta grille.",
    "Réponds à TOUS les commentaires dans la première heure : l'algorithme pousse les Reels qui ont de la conversation.",
    "Quand quelqu'un commente « LIMO » ou « DÉMO », envoie le lien en message privé dans la journée.",
    "Les heures (12:15 / 18:30 / samedi 10:00) sont un point de départ : après 2 semaines, regarde tes statistiques Instagram (onglet Audience → heures d'activité) et ajuste.",
    "Film « Signatures » : uniquement avec l'accord écrit des clients qui apparaissent, même floutés.",
    "Avant de publier une ancienne vidéo, vérifie qu'elle dit bien « Le bras droit du conseiller immo » (et pas l'ancienne accroche).",
]
for i, t in enumerate(tips):
    c = g.cell(row=3 + i, column=1, value=f"{i + 1}. {t}"); c.font = B; c.alignment = WR
g.column_dimensions["A"].width = 120
for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.save("CALENDRIER-PUBLICATION-LIMO.xlsx")
print("ok", k, "reels")
